import io
import uuid

import pytest
from docx import Document as DocxDocument

from app.services import attachment_service as svc

PNG = b"\x89PNG\r\n\x1a\n" + b"\x00" * 32
USER = uuid.uuid4()


def _docx(*paragraphs: str) -> bytes:
    doc = DocxDocument()
    doc.add_heading("Ưu đãi Q4", level=1)
    for p in paragraphs:
        doc.add_paragraph(p)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


async def test_document_uses_rag_parsers():
    att = await svc.process("brief.docx", _docx("Giảm 20% gói Business, mã QUYETTOAN26."), USER)
    assert (att.kind, att.mime_type, att.truncated) == ("document", svc.DOC_MIME["DOCX"], False)
    assert att.text.startswith("# Ưu đãi Q4") and "QUYETTOAN26" in att.text

    att = await svc.process("note.md", "# Ghi chú\r\n\r\n\r\n\r\nNội dung".encode(), USER)
    assert att.text == "# Ghi chú\n\nNội dung"


async def test_long_document_is_truncated():
    att = await svc.process("long.txt", (("a" * 100 + "\n") * 500).encode(), USER)
    assert att.truncated and len(att.text) <= svc.MAX_ATTACHMENT_CHARS + 2 and att.text.endswith("…")


async def test_image_goes_to_vision(monkeypatch):
    calls = []

    async def fake_describe(data, mime, user_id, thread_id):
        calls.append((mime, user_id, thread_id))
        return "**Chủ đề chính:** ly cà phê"

    monkeypatch.setattr(svc, "describe_image", fake_describe)
    att = await svc.process("../../anh.PNG", PNG, USER, "t1")
    assert (att.kind, att.filename, att.text) == ("image", "anh.PNG", "**Chủ đề chính:** ly cà phê")
    assert calls == [("image/png", USER, "t1")]


@pytest.mark.parametrize(
    "name,data,error",
    [
        ("virus.exe", b"MZ....", "Chỉ hỗ trợ"),
        ("fake.png", b"GIF89a....", "không phải ảnh"),
        ("fake.webp", b"RIFF\x00\x00\x00\x00WAVE", "không phải ảnh"),
        ("fake.pdf", b"hello", "không phải PDF"),
        ("empty.txt", b"   \n ", "Không trích được"),
        ("broken.docx", b"PK\x03\x04garbage", "Không đọc được"),
    ],
)
async def test_invalid_files_rejected(name, data, error):
    with pytest.raises(svc.AttachmentError, match=error):
        await svc.process(name, data, USER)


def test_format_for_context():
    assert svc.format_for_context("image", "a.png", "mô tả") == (
        "[Tệp người dùng đính kèm: a.png (ảnh – mô tả do AI phân tích)]\nmô tả"
    )
    assert svc.format_for_context("document", "b.pdf", "x", truncated=True).endswith("x\n(đã cắt bớt phần cuối)")
