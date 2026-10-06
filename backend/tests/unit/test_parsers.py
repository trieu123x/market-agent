import io

import pymupdf
from docx import Document as DocxDocument

from app.rag.parsers import (
    docx_to_markdown,
    html_to_markdown,
    parse_to_markdown,
    pdf_to_markdown,
    strip_table_of_contents,
    txt_to_markdown,
)


def test_pdf_to_markdown():
    doc = pymupdf.open()
    doc.new_page().insert_text((72, 72), "Brand guideline page one")
    doc.new_page().insert_text((72, 72), "Second page content")
    md = pdf_to_markdown(doc.tobytes())
    assert "Brand guideline page one" in md and "Second page content" in md


def test_docx_to_markdown_headings_lists_tables():
    d = DocxDocument()
    d.add_heading("Chiến dịch Tết", level=1)
    d.add_paragraph("Mục tiêu tăng nhận diện.")
    d.add_paragraph("Kênh Facebook", style="List Bullet")
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text = "KPI"
    t.rows[0].cells[1].text = "10%"
    buf = io.BytesIO()
    d.save(buf)
    md = docx_to_markdown(buf.getvalue())
    assert "# Chiến dịch Tết" in md
    assert "- Kênh Facebook" in md
    assert "| KPI | 10% |" in md


def test_txt_utf8_bom():
    assert txt_to_markdown("﻿Xin chào\r\n\r\n\r\n\r\nthế giới".encode()) == "Xin chào\n\nthế giới"


def test_html_strips_boilerplate():
    html = "<html><body><nav>Menu</nav><article><h2>Tiêu đề</h2><p>Nội dung <b>chính</b></p>" \
           "<script>alert(1)</script></article><footer>©</footer></body></html>"
    md = html_to_markdown(html)
    assert md == "## Tiêu đề\n\nNội dung chính"


def test_url_pdf_content_dispatch():
    doc = pymupdf.open()
    doc.new_page().insert_text((72, 72), "Remote pdf")
    assert "Remote pdf" in parse_to_markdown("URL", doc.tobytes(), "application/pdf")


def test_docx_word_toc_removed():
    d = DocxDocument()
    d.add_paragraph("NỘI QUY LAO ĐỘNG")
    d.add_paragraph("MỤC LỤC")
    d.add_paragraph("CHƯƠNG I: Quy định chung\t3")
    d.add_paragraph("CHƯƠNG II: Thời giờ làm việc\t4")
    d.add_heading("CHƯƠNG I: QUY ĐỊNH CHUNG", level=1)
    d.add_paragraph("Nội quy áp dụng cho toàn bộ người lao động.")
    buf = io.BytesIO()
    d.save(buf)
    md = parse_to_markdown("DOCX", buf.getvalue())
    assert md == "NỘI QUY LAO ĐỘNG\n\n# CHƯƠNG I: QUY ĐỊNH CHUNG\n\nNội quy áp dụng cho toàn bộ người lao động."


def test_pdf_style_toc_with_leaders_wrapped_lines_and_page_footer():
    md = (
        "Mục lục\nTrang\n"
        "Chương I: Quy định chung ................ 3\n"
        "Chương VII: Tạm thời chuyển người lao động làm công việc khác so với\n"
        "hợp đồng lao động ...................... 12\n"
        "\n2\n\n"
        "Chương XII: Điều khoản thi hành … 16\n"
        "Điều 5. Thời giờ làm việc không quá 08 giờ trong 01 ngày và không quá 48\n"
        "giờ trong 01 tuần."
    )
    assert strip_table_of_contents(md) == (
        "Điều 5. Thời giờ làm việc không quá 08 giờ trong 01 ngày và không quá 48\ngiờ trong 01 tuần."
    )


def test_untitled_leader_run_removed_only_when_long_enough():
    toc = "Giới thiệu ........ 1\nSản phẩm ........ 2\nGiá ........ 5\n\nNội dung chính."
    assert strip_table_of_contents(toc) == "Nội dung chính."
    short = "Giới thiệu ........ 1\nSản phẩm ........ 2\n\nNội dung chính."
    assert strip_table_of_contents(short) == short


def test_toc_title_without_entries_or_prose_is_kept():
    md = "## Contents\n\nBài viết này nói về chiến lược nội dung cho mùa Tết năm 2026\n\n## Kênh\n\nFacebook"
    assert strip_table_of_contents(md) == md
    md = "Mục lục sản phẩm mới gồm 3 dòng chính, ra mắt quý 4 năm 2026."
    assert strip_table_of_contents(md) == md


def test_markdown_upload_parsed_as_txt():
    from app.services.document_service import detect_file_type

    assert detect_file_type("Brief.MD", b"# Title") == "TXT"
